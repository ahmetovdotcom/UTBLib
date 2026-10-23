"""Receipt amounts by full catalogue departments (952$e), without allocation guesses."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from sqlalchemy import select
from .db import Acquisition, AcquisitionLine, Record, KsuEntry
from .ksu_auto import vals


def report(s, date_from, date_to):
    pairs=list(s.execute(select(Acquisition,KsuEntry).outerjoin(KsuEntry,KsuEntry.acquisition_id==Acquisition.id)
        .where(Acquisition.status=='posted',Acquisition.received_on>=date_from,Acquisition.received_on<=date_to)
        .order_by(Acquisition.received_on,Acquisition.id)))
    ids=[a.id for a,e in pairs];lines={}
    if ids:
        for line,record in s.execute(select(AcquisitionLine,Record).join(Record,Record.id==AcquisitionLine.record_id)
            .where(AcquisitionLine.acquisition_id.in_(ids),AcquisitionLine.quantity>0).order_by(AcquisitionLine.id)):
            lines.setdefault(line.acquisition_id,[]).append((line,record))
    columns={};rows=[];totals={};overall={'titles':0,'copies':0,'amount':Decimal(0),'unpriced_count':0}
    def bucket():return {'titles':set(),'copies':0,'amount':Decimal(0),'unpriced_count':0}
    def output(b):return {**b,'titles':len(b['titles']) if isinstance(b['titles'],set) else b['titles'],'amount':format(b['amount'],'.2f')}
    for a,e in pairs:
        groups={};total=bucket();issues=[]
        for line,record in lines.get(a.id,[]):
            departments=sorted(vals(record.marc,'952','e'));messages=[]
            if len(departments)==1:
                key='department:'+departments[0];label=departments[0]
            elif not departments:
                key='missing';label='Не указана кафедра';messages.append('Заполните 952$e «Кафедра» в карточке книги.')
            else:
                key='multiple';label='Несколько кафедр — не распределено'
                messages.append('Указано несколько кафедр: '+ '; '.join(departments)+'. Экземпляры и сумма учтены один раз в отдельной колонке, без распределения.')
            columns[key]=label
            known=line.unit_price is not None and line.currency=='KZT'
            if not known:messages.append('Не указана цена в тенге: исправьте поступление.')
            amount=line.unit_price*line.quantity if known else Decimal(0)
            for b in (groups.setdefault(key,bucket()),total):
                b['titles'].add(record.id);b['copies']+=line.quantity;b['amount']+=amount
                if not known:b['unpriced_count']+=line.quantity
            if messages:issues.append({'record_id':record.id,'title':' '.join(sorted(vals(record.marc,'245','a'))) or line.title_snapshot,'messages':messages})
        serialized={key:output(b) for key,b in groups.items()};row_total=output(total)
        for key,b in serialized.items():
            t=totals.setdefault(key,{'titles':0,'copies':0,'amount':Decimal(0),'unpriced_count':0})
            for metric in t:t[metric]+=Decimal(b[metric]) if metric=='amount' else b[metric]
        for metric in overall:overall[metric]+=Decimal(row_total[metric]) if metric=='amount' else row_total[metric]
        rows.append({'acquisition_id':a.id,'number':e.number if e else None,'year':e.year if e else None,
            'received_on':a.received_on.isoformat(),'source':a.source,'invoice':a.invoice,
            'groups':serialized,'total':row_total,'issues':issues})
    keys=sorted(columns,key=lambda key:(key in ('missing','multiple'),columns[key].casefold()))
    return {'date_from':date_from.isoformat(),'date_to':date_to.isoformat(),
        'columns':[{'key':key,'label':columns[key]} for key in keys],'rows':rows,
        'totals':{key:output(totals[key]) for key in keys},'total':output(overall),
        'issue_count':sum(len(r['issues']) for r in rows)}


def export_xlsx(data):
    # Same server-side XLSX library as the existing first-sheet exporter.
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    w=Workbook();ws=w.active;ws.title='По кафедрам'
    def put(row,col,value):
        c=ws.cell(row,col,value)
        if isinstance(value,str):c.data_type='s'
        return c
    groups=data['columns']+[{'key':'total','label':'Всего'}];width=4+3*len(groups)
    put(1,1,f"Поступления по кафедрам: {data['date_from']} — {data['date_to']}. Суммы в тенге, без доставки.")
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width)
    for i,label in enumerate(['№ КСУ','Дата поступления','Источник поступления','Документ'],1):
        put(2,i,label);ws.merge_cells(start_row=2,start_column=i,end_row=3,end_column=i)
    for i,g in enumerate(groups):
        col=5+i*3;put(2,col,g['label']);ws.merge_cells(start_row=2,start_column=col,end_row=2,end_column=col+2)
        for j,label in enumerate(['Наименования','Экземпляры','Сумма, ₸']):put(3,col+j,label)
    for n,r in enumerate(data['rows'],4):
        values=[f"{r['number']}/{r['year']}" if r['number'] is not None else 'Без КСУ',date.fromisoformat(r['received_on']),r['source'],r['invoice']]
        for i,v in enumerate(values,1):put(n,i,v)
        for i,g in enumerate(groups):
            b=r['total'] if g['key']=='total' else r['groups'].get(g['key'],{})
            for j,k in enumerate(['titles','copies','amount']):put(n,5+3*i+j,Decimal(b.get(k,'0')) if k=='amount' else b.get(k,0))
    last=4+len(data['rows']);put(last,1,'Итого')
    for i,g in enumerate(groups):
        b=data['total'] if g['key']=='total' else data['totals'][g['key']]
        for j,k in enumerate(['titles','copies','amount']):put(last,5+3*i+j,Decimal(b[k]) if k=='amount' else b[k])
    for row in ws:
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='center')
            if c.row<=3 or c.row==last:c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='E8EEF8')
            if c.column>=5 and (c.column-5)%3==2:c.number_format='#,##0.00'
            if isinstance(c.value,date):c.number_format='DD.MM.YYYY'
    for i in range(1,width+1):ws.column_dimensions[get_column_letter(i)].width=18 if i not in (3,4) else 30
    ws.row_dimensions[2].height=80;ws.row_dimensions[1].height=35
    ws.freeze_panes='E4';ws.print_title_rows='1:3';ws.page_setup.orientation='landscape'
    ws.print_area=f'A1:{get_column_letter(width)}{last}'
    notes=w.create_sheet('Пояснения')
    for text in ['Кафедры — полные значения 952$e из текущих карточек каталога.',
        'Каждый экземпляр и его стоимость учитываются один раз. Несколько кафедр не распределяются автоматически.',
        'Наименования считаются отдельно в каждой партии. Общий итог — сумма по партиям, не уникальные названия фонда.',
        'Отменённые поступления и экземпляры без поступления исключены. Списания не уменьшают исторические поступления.',
        'Период определяется по дате поступления включительно. Суммы в тенге, без доставки.']:
        notes.append([text])
    for r in data['rows']:
        for issue in r['issues']:
            notes.append([f"Поступление №{r['acquisition_id']}",f"Запись №{issue['record_id']}",issue['title'],' '.join(issue['messages'])])
    for row in notes:
        for c in row:
            if isinstance(c.value,str):c.data_type='s'
            c.alignment=Alignment(wrap_text=True,vertical='top')
    for col in ('A','B','C','D'):notes.column_dimensions[col].width=45
    out=BytesIO();w.save(out);return out.getvalue()
