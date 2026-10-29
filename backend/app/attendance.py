"""Attendance: read-only aggregates with shared filters and local calendar boundaries."""
import csv
import io
from datetime import date, datetime, timedelta, timezone
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator
from sqlalchemy import Integer, String, case, cast, func, or_, select
from .db import User
from .reader_core import CATEGORIES, STATUSES, TZ, local_bounds
from .reader_models import Branch, LibraryVisit as Visit, Reader

BUCKETS = {
    'bachelor': 'Бакалавры (БАК)', 'master': 'Магистранты (МАГ)',
    'faculty': 'Преподаватели / ППС (ПОҚ)', 'staff': 'Сотрудники',
    'external': 'Внешние пользователи', 'doctoral': 'Докторанты',
    'other': 'Прочие / категория не уточнена',
}
SOURCES = {'desk': 'Рабочее место', 'kiosk': 'Киоск', 'turnstile': 'Турникет'}
DIMENSIONS = {
    'faculty': 'Факультет', 'department': 'Кафедра', 'education_program': 'Образовательная программа',
    'study_group': 'Учебная группа', 'course': 'Курс', 'cohort_year': 'Год поступления',
    'branch_id': 'Пункт обслуживания', 'scanned_by': 'Отметил', 'source': 'Точка отметки',
    'category': 'Категория читателя', 'status': 'Статус читателя',
}
READER_FILTERS = ('category', 'status', 'faculty', 'department', 'education_program', 'study_group', 'course', 'cohort_year')
LIST_FILTERS = (*READER_FILTERS, 'bucket', 'branch_id', 'scanned_by', 'source', 'weekday')
FILTER_LABELS = {**DIMENSIONS, 'bucket': 'Группа отчёта', 'weekday': 'Дни недели',
                 'q': 'Поиск читателя', 'time_from': 'Время с', 'time_to': 'Время до'}
WEEKDAYS = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс']
RULES = ('Посещения — зарегистрированные визиты, уникальные читатели — разные читатели за выбранный период. '
         'Уникальные читатели в итогах пересчитываются, а не складываются по дням. '
         'Повторные сканы, уже исключённые при регистрации, не учитываются. '
         'Категория и учебные данные берутся из текущих карточек читателей; исторических снимков этих полей нет.')


class Filters(BaseModel):
    model_config = ConfigDict(extra='forbid')
    date_from: date = Field(default_factory=lambda: datetime.now(TZ).date().replace(day=1))
    date_to: date = Field(default_factory=lambda: datetime.now(TZ).date())
    group_by: Literal['day', 'week', 'month'] = 'day'
    breakdown: Literal['faculty', 'department', 'education_program', 'study_group', 'course',
                       'cohort_year', 'branch_id', 'scanned_by', 'source', 'category', 'status'] = 'faculty'
    bucket: list[str] = Field(default_factory=list, max_length=20)
    category: list[str] = Field(default_factory=list, max_length=30)
    status: list[str] = Field(default_factory=list, max_length=10)
    faculty: list[str] = Field(default_factory=list, max_length=100)
    department: list[str] = Field(default_factory=list, max_length=100)
    education_program: list[str] = Field(default_factory=list, max_length=100)
    study_group: list[str] = Field(default_factory=list, max_length=100)
    course: list[str] = Field(default_factory=list, max_length=30)
    cohort_year: list[str] = Field(default_factory=list, max_length=100)
    branch_id: list[int] = Field(default_factory=list, max_length=100)
    scanned_by: list[int] = Field(default_factory=list, max_length=100)
    source: list[Literal['desk', 'kiosk', 'turnstile']] = Field(default_factory=list, max_length=3)
    weekday: list[int] = Field(default_factory=list, max_length=7)
    time_from: str = ''
    time_to: str = ''
    q: str = Field(default='', max_length=200)

    @field_validator('time_from', 'time_to')
    @classmethod
    def valid_time(cls, value):
        import re
        if value and not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d', value):
            raise ValueError('Время должно быть в формате ЧЧ:ММ')
        return value

    @model_validator(mode='after')
    def valid_filters(self):
        if not date(1900, 1, 1) <= self.date_from <= self.date_to <= date(2200, 12, 31):
            raise ValueError('Укажите корректный период с 1900 по 2200 год')
        if (self.date_to - self.date_from).days >= 3660:
            raise ValueError('Максимальный период одного отчёта — 3660 дней')
        if any(v not in BUCKETS for v in self.bucket):
            raise ValueError('Неизвестная группа отчёта')
        if any(v not in range(7) for v in self.weekday):
            raise ValueError('Дни недели должны быть от 0 (понедельник) до 6 (воскресенье)')
        if any(v < 1 for v in (*self.branch_id, *self.scanned_by)):
            raise ValueError('Некорректный идентификатор фильтра')
        for key in READER_FILTERS:
            if any(len(v) > 300 for v in getattr(self, key)):
                raise ValueError('Слишком длинное значение фильтра')
        self.q = self.q.strip()
        return self


def parse_filters(request: Request):
    values = {key: request.query_params.getlist(key) if key in LIST_FILTERS else request.query_params[key]
              for key in Filters.model_fields if key in request.query_params}
    if set(request.query_params) - set(Filters.model_fields) - {'offset', 'limit', 'format'}:
        raise HTTPException(422, 'Неизвестный параметр отчёта')
    try:
        return Filters.model_validate(values)
    except ValidationError as error:
        raise HTTPException(422, '; '.join(e['msg'].removeprefix('Value error, ') for e in error.errors())) from None


def bucket_expression():
    # Every reader belongs to one bucket. Unspecified students are not assumed to be bachelors.
    return case(
        (or_(Reader.category.in_(['teacher', 'faculty']), Reader.is_teaching_staff.is_(True)), 'faculty'),
        *[(Reader.category == key, key) for key in ('bachelor', 'master', 'staff', 'external', 'doctoral')],
        else_='other',
    )


def local_expressions(s):
    if s.bind.dialect.name == 'sqlite':
        # Same IANA timezone as PostgreSQL, including Kazakhstan's 2024 offset change.
        def local(value):
            if value is None: return None
            stamp = datetime.fromisoformat(value)
            if stamp.tzinfo is None: stamp = stamp.replace(tzinfo=timezone.utc)
            return stamp.astimezone(TZ).strftime('%Y-%m-%d %H:%M:%S')
        s.connection().connection.driver_connection.create_function('attendance_local', 1, local, deterministic=True)
        stamp = func.attendance_local(Visit.scanned_at)
        day = func.substr(stamp, 1, 10)
        hour = cast(func.strftime('%H', stamp), Integer)
        minute = hour * 60 + cast(func.strftime('%M', stamp), Integer)
        weekday = (cast(func.strftime('%w', stamp), Integer) + 6) % 7
        week = func.date(stamp, 'weekday 0', '-6 days')
        month = func.substr(stamp, 1, 7) + '-01'
    else:
        stamp = func.timezone(str(TZ), Visit.scanned_at)
        day = func.to_char(stamp, 'YYYY-MM-DD')
        hour = cast(func.extract('hour', stamp), Integer)
        minute = hour * 60 + cast(func.extract('minute', stamp), Integer)
        weekday = cast(func.extract('isodow', stamp), Integer) - 1
        week = func.to_char(func.date_trunc('week', stamp), 'YYYY-MM-DD')
        month = func.to_char(stamp, 'YYYY-MM') + '-01'
    return {'day': day, 'week': week, 'month': month, 'hour': hour, 'minute': minute, 'weekday': weekday}


def query_context(s, f):
    local, bucket = local_expressions(s), bucket_expression()
    start, _ = local_bounds(f.date_from); _, end = local_bounds(f.date_to)
    conditions = [Visit.scanned_at >= start, Visit.scanned_at < end]
    for key in READER_FILTERS:
        if getattr(f, key):
            conditions.append(func.coalesce(cast(getattr(Reader, key), String), '').in_(getattr(f, key)))
    for key in ('branch_id', 'scanned_by', 'source'):
        if getattr(f, key): conditions.append(getattr(Visit, key).in_(getattr(f, key)))
    if f.bucket: conditions.append(bucket.in_(f.bucket))
    if f.weekday: conditions.append(local['weekday'].in_(f.weekday))
    def minutes(value):
        hour, minute = map(int, value.split(':')); return hour * 60 + minute
    if f.time_from and f.time_to and f.time_from > f.time_to:
        conditions.append(or_(local['minute'] >= minutes(f.time_from), local['minute'] <= minutes(f.time_to)))
    else:
        if f.time_from: conditions.append(local['minute'] >= minutes(f.time_from))
        if f.time_to: conditions.append(local['minute'] <= minutes(f.time_to))
    if f.q:
        pattern = '%' + f.q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
        conditions.append(or_(*(c.ilike(pattern, escape='\\') for c in (Reader.full_name, Reader.ticket_number, Reader.barcode))))
    def query(*columns):
        return select(*columns).select_from(Visit).join(Reader, Reader.id == Visit.reader_id).where(*conditions)
    return query, local, bucket


def counts():
    return func.count(Visit.id), func.count(func.distinct(Visit.reader_id))


def blank_row():
    return {'visits': 0, 'unique': 0, 'categories': {key: {'visits': 0, 'unique': 0} for key in BUCKETS}}


def grouped(s, query, dimension, bucket):
    result = {}
    for period, key, visits, unique in s.execute(query(dimension, bucket, *counts()).group_by(dimension, bucket)):
        row = result.setdefault(str(period), blank_row())
        row['categories'][key] = {'visits': visits, 'unique': unique}
        row['visits'] += visits
        row['unique'] += unique  # Mutually exclusive, current reader categories.
    return result


def report(s, f):
    query, local, bucket = query_context(s, f)
    totals = blank_row()
    for key, visits, unique in s.execute(query(bucket, *counts()).group_by(bucket)):
        totals['categories'][key] = {'visits': visits, 'unique': unique}
        totals['visits'] += visits; totals['unique'] += unique
    periods = grouped(s, query, local[f.group_by], bucket)
    months = grouped(s, query, local['month'], bucket) if f.group_by == 'day' else {}
    dates = [f.date_from + timedelta(days=i) for i in range((f.date_to - f.date_from).days + 1)]
    dates = [d for d in dates if not f.weekday or d.weekday() in f.weekday]
    def period_key(d):
        if f.group_by == 'week': return d - timedelta(days=d.weekday())
        return d.replace(day=1) if f.group_by == 'month' else d
    ranges = {}
    for day in dates:
        key = str(period_key(day))
        if key not in ranges: ranges[key] = [day, day]
        else: ranges[key][1] = day
    rows = [{'period': key, 'date_from': str(bounds[0]), 'date_to': str(bounds[1]),
             **periods.get(key, blank_row())} for key, bounds in ranges.items()]
    dimension = getattr(Reader, f.breakdown) if f.breakdown in READER_FILTERS else getattr(Visit, f.breakdown)
    dimension = func.coalesce(cast(dimension, String), '')
    labels = {}
    if f.breakdown == 'branch_id':
        labels = {str(b.id): b.name for b in s.scalars(select(Branch))}
    elif f.breakdown == 'scanned_by':
        labels = {str(i): name for i, name in s.execute(select(User.id, User.display_name))}
    elif f.breakdown in ('source', 'category', 'status'):
        labels = {'source': SOURCES, 'category': CATEGORIES, 'status': STATUSES}[f.breakdown]
    breakdown = [{'key': str(key), 'label': labels.get(str(key), str(key)) or 'Не указано', 'visits': v, 'unique': u}
                 for key, v, u in s.execute(query(dimension, *counts()).group_by(dimension))]
    breakdown.sort(key=lambda row: (-row['visits'], row['label']))
    daily = list(s.execute(query(local['day'], func.count(Visit.id)).group_by(local['day'])))
    peak = max(daily, key=lambda r: (r[1], r[0]), default=None)
    hours = {h: (v, u) for h, v, u in s.execute(query(local['hour'], *counts()).group_by(local['hour']))}
    weekdays = {d: (v, u) for d, v, u in s.execute(query(local['weekday'], *counts()).group_by(local['weekday']))}
    return {
        'filters': f.model_dump(mode='json'), 'timezone': str(TZ), 'rules': RULES,
        'columns': [{'key': key, 'label': label} for key, label in BUCKETS.items()
                    if key not in ('doctoral', 'other') or totals['categories'][key]['visits']],
        'totals': totals, 'rows': rows,
        'months': [{'period': key, **months.get(key, blank_row())} for key in sorted({str(d.replace(day=1)) for d in dates})] if f.group_by == 'day' else [],
        'summary': {'selected_days': len(dates), 'active_days': len(daily),
                    'average_per_day': round(totals['visits'] / len(dates), 2) if dates else 0,
                    'visits_per_reader': round(totals['visits'] / totals['unique'], 2) if totals['unique'] else 0,
                    'peak': {'date': peak[0], 'visits': peak[1]} if peak else None},
        'breakdown': breakdown,
        'hours': [{'hour': h, 'visits': hours.get(h, (0, 0))[0], 'unique': hours.get(h, (0, 0))[1]} for h in range(24)],
        'weekdays': [{'day': d, 'label': label, 'visits': weekdays.get(d, (0, 0))[0], 'unique': weekdays.get(d, (0, 0))[1]} for d, label in enumerate(WEEKDAYS)],
    }


def options(s):
    result = {}
    for key in READER_FILTERS:
        values = s.scalars(select(func.coalesce(cast(getattr(Reader, key), String), '')).distinct()).all()
        labels = CATEGORIES if key == 'category' else STATUSES if key == 'status' else {}
        values = set(values) | (set(labels) if labels else {''})
        result[key] = [{'value': v, 'label': labels.get(v, v) or 'Не указано'} for v in sorted(values)]
    result['bucket'] = [{'value': k, 'label': v} for k, v in BUCKETS.items()]
    result['branch_id'] = [{'value': str(b.id), 'label': b.name + ('' if b.active else ' (закрыт)')} for b in s.scalars(select(Branch).order_by(Branch.position, Branch.name))]
    result['scanned_by'] = [{'value': str(i), 'label': name} for i, name in s.execute(select(User.id, User.display_name).where(User.id.in_(select(Visit.scanned_by))).order_by(User.display_name))]
    result['source'] = [{'value': k, 'label': v} for k, v in SOURCES.items()]
    result['weekday'] = [{'value': str(i), 'label': v} for i, v in enumerate(WEEKDAYS)]
    return {'filters': result, 'dimensions': DIMENSIONS, 'timezone': str(TZ)}


def export_report(data, f, format, filter_options=None):
    from .readers import safe_cell
    keys = [c['key'] for c in data['columns']]
    def values(row):
        return [row['categories'][key]['visits'] for key in keys] + [row['visits'], row['unique']]
    table = [['Отчёт по посещаемости'], [f'{f.date_from} — {f.date_to} · {TZ}'],
             ['Период', *[c['label'] for c in data['columns']], 'Всего посещений', 'Уникальных читателей']]
    month_lookup = {m['period'][:7]: m for m in data['months']}
    for i, row in enumerate(data['rows']):
        label = row['date_from'] if row['date_from'] == row['date_to'] else row['date_from'] + ' — ' + row['date_to']
        table.append([label, *values(row)])
        if f.group_by == 'day' and (i + 1 == len(data['rows']) or data['rows'][i + 1]['period'][:7] != row['period'][:7]):
            table.append(['Итого за ' + row['period'][:7], *values(month_lookup[row['period'][:7]])])
    table.extend([['Итого за период', *values(data['totals'])], [], [RULES],
                  ['Разбивка: ' + DIMENSIONS[f.breakdown], 'Посещений', 'Уникальных читателей']])
    table.extend([[r['label'], r['visits'], r['unique']] for r in data['breakdown']])
    table.extend([[], ['Час (Алматы)', 'Посещений', 'Уникальных читателей']])
    table.extend([[f"{r['hour']:02d}:00", r['visits'], r['unique']] for r in data['hours']])
    table.extend([[], ['День недели', 'Посещений', 'Уникальных читателей']])
    table.extend([[r['label'], r['visits'], r['unique']] for r in data['weekdays']])
    table.extend([[], ['Применённые фильтры']])
    for key, value in f.model_dump(mode='json').items():
        if key in FILTER_LABELS and value != [] and value != '':
            names = {str(o['value']): o['label'] for o in (filter_options or {}).get(key, [])}
            display = ', '.join(names.get(str(v), str(v)) or 'Не указано' for v in value) if isinstance(value, list) else value
            table.append([FILTER_LABELS[key], display])
    filename = f'attendance-{f.date_from}-{f.date_to}.{format}'
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    if format == 'csv':
        out = io.StringIO()
        csv.writer(out, delimiter=';').writerows([[safe_cell(v) for v in row] for row in table])
        return Response('\ufeff' + out.getvalue(), media_type='text/csv; charset=utf-8', headers=headers)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Посещаемость'
    for row in table:
        ws.append([safe_cell(v) if isinstance(v, str) else v for v in row])
    width = len(keys) + 3
    for row in ws:
        for cell in row: cell.alignment = Alignment(vertical='center', wrap_text=True)
        if row[0].row == 3 or str(row[0].value).startswith('Итого'):
            for cell in row[:width]:
                cell.font = Font(bold=True, color='052744'); cell.fill = PatternFill('solid', fgColor='E4EEF5')
    ws.column_dimensions['A'].width = 33
    for index in range(2, width + 1): ws.column_dimensions[get_column_letter(index)].width = 21
    ws.row_dimensions[3].height = 48; ws.freeze_panes = 'B4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = '1:3'
    out = io.BytesIO(); wb.save(out)
    return Response(out.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers=headers)


def make_router(auth, db):
    router = APIRouter(prefix='/api/statistics/attendance', dependencies=[Depends(auth)])

    @router.get('/options')
    def get_options(s=Depends(db)):
        return options(s)

    @router.get('')
    def get_report(f=Depends(parse_filters), s=Depends(db)):
        return report(s, f)

    @router.get('/export')
    def export(format: Literal['xlsx', 'csv'] = 'xlsx', f=Depends(parse_filters), s=Depends(db)):
        return export_report(report(s, f), f, format, options(s)['filters'])

    @router.get('/journal')
    def journal(f=Depends(parse_filters), offset: int = Query(0, ge=0),
                limit: int = Query(50, ge=1, le=100), s=Depends(db)):
        query, _, bucket = query_context(s, f)
        total = s.scalar(query(func.count(Visit.id)))
        stmt = query(Visit.id, Visit.reader_id, Visit.scanned_at, Reader.full_name, Reader.ticket_number,
                     Reader.category, Reader.faculty, Reader.study_group, Branch.name.label('branch'),
                     User.display_name.label('actor'), Visit.source, bucket.label('bucket'))
        stmt = stmt.join(Branch, Branch.id == Visit.branch_id).join(User, User.id == Visit.scanned_by)
        items = []
        for row in s.execute(stmt.order_by(Visit.scanned_at.desc(), Visit.id.desc()).offset(offset).limit(limit)):
            item = dict(row._mapping); stamp = item['scanned_at']
            if stamp.tzinfo is None: stamp = stamp.replace(tzinfo=timezone.utc)
            item['scanned_at'] = stamp.isoformat(); items.append(item)
        return {'items': items, 'total': total}

    return router
