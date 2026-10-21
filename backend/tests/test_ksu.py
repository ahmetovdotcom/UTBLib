from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from test_catalog import client,make_type,payload,field
from test_acquisitions import add_body,save
from app.ksu_auto import classify,DEFAULT_RULES,norm


def receipt(c,code='КУК',language='kaz',price='10.10'):
    t=make_type(c);p=payload(t)
    p['marc']['fields'] += [field('952','a',code),field('041','a',language),field('080','a','3')]
    r=save(c,p);body=add_body(r['id'],3,price);body['document_on']='2026-10-20'
    body['lines'].append({**body['lines'][0],'quantity':2})
    response=c.post('/api/acquisitions',json=body);assert response.status_code==200,response.text
    return response.json(),r,body

def report(c):return c.get('/api/ksu/report?year=2026').json()

def test_automatic_registration_no_manual_form_and_no_duplicates(client):
    a,r,b=receipt(client)
    data=report(client)
    assert data['issue_count']==0 and len(data['rows'])==1
    assert data['totals']['titles']==1 and data['totals']['copies']==5
    assert data['totals']['on_amount']=='50.50' and data['totals']['off_copies']==0
    assert data['totals']['study_kaz']==data['totals']['sector_3']==5
    assert data['rows'][0]['document_on']=='2026-10-20'
    assert client.post('/api/acquisitions',json=b).json()['id']==a['id']
    assert report(client)==data
    assert client.get('/api/records/'+str(r['id'])+'/copies').json()['total']==5
    assert client.put('/api/ksu/entries/'+str(a['id']),json={}).status_code==409
    response=client.get('/api/ksu/export?year=2026');assert response.status_code==200,response.text
    ws=load_workbook(BytesIO(response.content)).worksheets[0]
    assert ws['E4'].value==1 and ws['F4'].value==5 and ws['G4'].value==50.5
    assert ws['H4'].value==0 and ws['D4'].value.day==20


def test_classification_conflict_is_not_double_counted_and_correcting_catalog_updates_report(client):
    a,r,b=receipt(client,'КУК','rus')
    data=report(client);assert data['issue_count']>0
    assert data['totals']['copies']==5 and data['totals']['kind_total']==0
    assert client.get('/api/ksu/export?year=2026').status_code==409
    p={k:r[k] for k in ('type_id','version','marc')}
    for f in p['marc']['fields']:
        if f['tag']=='041':f['subfields'][0]['value']='kaz'
    response=client.put('/api/records/'+str(r['id']),json=p);assert response.status_code==200,response.text
    assert report(client)['issue_count']==0 and report(client)['totals']['study_kaz']==5


def test_custom_rules_reference_and_media(client):
    a,r,b=receipt(client,'КСК')
    assert report(client)['issue_count']>0
    rules=client.get('/api/ksu/rules').json()['items'];rule=next(x for x in rules if x['code']=='КСК')
    rule['kind']='science'
    response=client.put('/api/ksu/rules',json=rule);assert response.status_code==200,response.text
    assert client.put('/api/ksu/rules',json=rule).status_code==409
    assert report(client)['totals']['science_kaz']==5 and report(client)['issue_count']==0
    mapping={norm(k):v for k,v in DEFAULT_RULES.items()}
    for code,kind,lang in [('КНИ','science','foreign'),('КНК','science','kaz'),('КНР','science','rus'),('КУИ','study','foreign'),('КУМК','study','kaz'),('КУМР','study','rus'),('КХК','fiction','kaz'),('КХР','fiction','rus'),('КХИ','fiction','foreign')]:
        c,issues=classify({'fields':[field('952','a',code),field('080','a','8')]},mapping)
        assert not issues and c['kind']==kind and c['language']==lang,(code,issues)
    c,issues=classify({'fields':[field('952','a','CD/каз.яз.'),field('080','a','0')]},mapping)
    assert c['media']=='cd' and c['language']=='kaz' and issues # no invented science/study category


def test_receipt_correction_cancel_and_existing_fund(client):
    a,r,b=receipt(client)
    client.post('/api/records/'+str(r['id'])+'/copies',json={'request_key':str(uuid4()),'unit_price':'500'})
    number=report(client)['rows'][0]['number']
    update={k:a[k] for k in ('version','invoice','received_on','document_on','source','note')}
    update.update(request_key=str(uuid4()),reason='Цена исправлена',lines=[{k:l[k] for k in ('id','record_id','quantity','unit_price','currency','sigla','note')} for l in a['lines']])
    update['lines'][0]['unit_price']='20'
    response=client.put('/api/acquisitions/'+str(a['id']),json=update);assert response.status_code==200,response.text
    assert report(client)['totals']['on_amount']=='80.20' and report(client)['totals']['copies']==5
    assert report(client)['rows'][0]['number']==number
    response=client.post('/api/acquisitions/'+str(a['id'])+'/cancel',json={'request_key':str(uuid4()),'version':2,'reason':'Ошибка'})
    assert response.status_code==200,response.text
    data=report(client)
    assert data['totals']['copies']==0 and data['rows']==[]
    assert data['issue_count']==0 and data['unpriced_count']==0
    assert client.get('/api/ksu/export?year=2026').status_code==200
    # Cancellation preserves the allocated number; the next receipt does not reuse it.
    new=client.post('/api/acquisitions',json=add_body(r['id'],1,'20'))
    assert new.status_code==200,new.text
    rows=report(client)['rows']
    assert len(rows)==1 and rows[0]['acquisition_id']==new.json()['id']
    assert rows[0]['number']==number+1
    assert report(client)['totals']['copies']==1


def test_unknown_price_is_not_zero_and_empty_document_date_not_invented(client):
    a,r,b=receipt(client,price=None)
    data=report(client);assert data['unpriced_count']==5 and data['issue_count']>0
    assert client.get('/api/ksu/export?year=2026').status_code==409
    other=add_body(r['id'],1,'0');other['received_on']='2027-01-01'
    client.post('/api/acquisitions',json=other)
    new=client.get('/api/ksu/report?year=2027').json()
    assert new['rows'][0]['number']==1 and new['rows'][0]['document_on'] is None
    assert new['unpriced_count']==0 and new['totals']['on_amount']=='0.00'
    assert client.get('/api/ksu/export?year=2027').status_code==200


def test_ksu_rights(client):
    a,r,b=receipt(client)
    from test_security import add,login
    from fastapi.testclient import TestClient
    from app.main import app
    add(client,username='ksuviewer',permissions=['ksu.view'],role='viewer')
    with TestClient(app) as c:
        login(c,'ksuviewer')
        assert c.get('/api/ksu/report?year=2026').status_code==200
        assert c.put('/api/ksu/rules',json={'code':'КУК','version':0,'kind':'science'}).status_code==403
    add(client,username='noview',permissions=[],role='viewer')
    with TestClient(app) as c:
        login(c,'noview')
        assert c.get('/api/ksu/report?year=2026').status_code==403


def test_multiple_952_values_conflict_and_same_category_not_duplicated():
    mapping={norm(k):v for k,v in DEFAULT_RULES.items()}
    c,issues=classify({'fields':[field('952','a','КУК'),field('952','a','КУР'),field('080','a','3')]},mapping)
    assert not c['language'] and issues
    c,issues=classify({'fields':[field('952','a','КУК'),field('952','a','КУМК'),field('080','a','3')]},mapping)
    assert c['language']=='kaz' and c['kind']=='study' and not issues
