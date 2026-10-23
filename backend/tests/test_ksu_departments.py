from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from test_catalog import client,make_type,payload,field
from test_acquisitions import add_body,save

URL='/api/ksu/departments/report?date_from=2026-10-21&date_to=2026-10-21'
EXPORT=URL.replace('/report?','/export?')

def book(c,t,departments):
    p=payload(t)
    p['marc']['fields'] += [field('952','e',v) for v in departments]
    return save(c,p)

def add(c,r,q=3,price='10.10',day='2026-10-21'):
    b=add_body(r['id'],q,price);b['received_on']=day
    response=c.post('/api/acquisitions',json=b);assert response.status_code==200,response.text
    return response.json()

def test_department_full_names_totals_dates_and_no_existing_fund(client):
    t=make_type(client);r=book(client,t,['Химия, химическая технология и экология'])
    b=add_body(r['id'],3,'10.10');b['lines'].append({**b['lines'][0],'quantity':2})
    assert client.post('/api/acquisitions',json=b).status_code==200
    add(client,r,1,'0');add(client,r,7,'8',day='2026-10-22')
    assert client.post(f"/api/records/{r['id']}/copies",json={'request_key':str(uuid4()),'unit_price':'999'}).status_code==201
    d=client.get(URL).json()
    assert len(d['columns'])==1 and d['columns'][0]['label']=='Химия, химическая технология и экология'
    assert d['total']=={'titles':2,'copies':6,'amount':'50.50','unpriced_count':0}
    assert d['totals'][d['columns'][0]['key']]==d['total']
    assert d['rows'][0]['total']['titles']==1 and d['issue_count']==0
    response=client.get(EXPORT);assert response.status_code==200,response.text
    ws=load_workbook(BytesIO(response.content)).worksheets[0]
    assert ws.title=='По кафедрам' and 'Поступления по кафедрам' in ws['A1'].value
    assert 'ksu-departments-' in response.headers['content-disposition']
    assert ws['E2'].value==d['columns'][0]['label']
    assert ws['E4'].value==1 and ws['F4'].value==5 and ws['G4'].value==50.5
    assert client.get(URL.replace('2026-10-21&','2026-10-22&')).status_code==422

def test_missing_multiple_departments_and_unknown_price(client):
    t=make_type(client)
    for specs in [[],['Экономика','Туризм'],['Туризм','Туризм']]:add(client,book(client,t,specs),2,'5')
    d=client.get(URL).json()
    assert set(d['totals'])=={'missing','multiple','department:Туризм'}
    assert {c['label'] for c in d['columns']}=={'Не указана кафедра','Несколько кафедр — не распределено','Туризм'}
    assert d['total']['copies']==sum(b['copies'] for b in d['totals'].values())==6
    assert d['total']['amount']=='30.00' and d['issue_count']==2
    # Explicitly unallocated columns may be exported, with explanations.
    assert client.get(EXPORT).status_code==200
    add(client,book(client,t,['Туризм']),1,None)
    d=client.get(URL).json();assert d['total']['unpriced_count']==1 and d['total']['amount']=='30.00'
    assert client.get(EXPORT).status_code==409

def test_catalog_changes_cancel_and_excel_text_safety(client):
    t=make_type(client);r=book(client,t,['=1+1']);a=add(client,r)
    response=client.get(EXPORT)
    ws=load_workbook(BytesIO(response.content)).worksheets[0]
    assert ws['E2'].value=='=1+1' and ws['E2'].data_type=='s'
    b={k:r[k] for k in ('type_id','version','marc')}
    b['marc']['fields'][-1]['subfields'][0]['value']='Полное новое название'
    changed=client.put(f"/api/records/{r['id']}",json=b);assert changed.status_code==200,changed.text
    assert client.get(URL).json()['columns'][0]['label']=='Полное новое название'
    cancelled=client.post(f"/api/acquisitions/{a['id']}/cancel",json={'request_key':str(uuid4()),'version':a['version'],'reason':'Ошибка'})
    assert cancelled.status_code==200,cancelled.text
    d=client.get(URL).json();assert d['rows']==[] and d['columns']==[] and d['total']['copies']==0
    assert client.get(EXPORT).status_code==200

def test_department_rights(client):
    from test_security import add as user,login
    from fastapi.testclient import TestClient
    from app.main import app
    user(client,username='specialtyviewer',permissions=['ksu.view'],role='viewer')
    user(client,username='specialtydenied',permissions=[],role='viewer')
    for username,status in [('specialtyviewer',200),('specialtydenied',403)]:
        with TestClient(app) as c:
            login(c,username)
            assert c.get(URL).status_code==status
            assert c.get(EXPORT).status_code==status

def test_department_groups_ignore_specialties_and_recalculate(client):
    t=make_type(client);records=[]
    for department,specialty in [('Финансы и учёт','6B04108 — Финансы'),('Финансы и учёт','6B04109 — Учёт и аудит'),(None,'6B04108 — Финансы')]:
        p=payload(t)
        p['marc']['fields'].append(field('952','j',specialty))
        if department:p['marc']['fields'].append(field('952','e',department))
        records.append(save(client,p))
    b=add_body(records[0]['id'],2,'15')
    b['lines'] += [{**b['lines'][0],'record_id':r['id']} for r in records[1:]]
    assert client.post('/api/acquisitions',json=b).status_code==200
    d=client.get(URL).json()
    assert d['totals']['department:Финансы и учёт']=={'titles':2,'copies':4,'amount':'60.00','unpriced_count':0}
    assert d['totals']['missing']['copies']==2
    assert d['total']=={'titles':3,'copies':6,'amount':'90.00','unpriced_count':0}
    assert '952$e' in d['rows'][0]['issues'][0]['messages'][0]
    assert client.get(URL.replace('/departments/','/specialties/')).json()==d
    r=records[0];b={k:r[k] for k in ('type_id','version','marc')}
    b['marc']['fields'][-1]['subfields'][0]['value']='Экономика и управление'
    changed=client.put(f"/api/records/{r['id']}",json=b)
    assert changed.status_code==200,changed.text
    updated=client.get(URL).json()
    assert updated['total']==d['total']
    assert updated['totals']['department:Финансы и учёт']['copies']==2
    assert updated['totals']['department:Экономика и управление']['copies']==2
