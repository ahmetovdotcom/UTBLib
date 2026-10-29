from datetime import datetime, timezone
import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from test_catalog import client
from test_security import add, login
from app.main import app, db
from app.db import User
from app.reader_core import TZ
from app.reader_models import Branch, Reader, LibraryVisit

PATH = '/api/statistics/attendance'
PERIOD = {'date_from': '2026-09-30', 'date_to': '2026-10-06'}

def seed(s):
    actor = s.scalar(select(User.id))
    b1 = Branch(code='A', name='Зал'); b2 = Branch(code='B', name='Абонемент', active=False)
    s.add_all([b1, b2]); s.flush()
    people = []
    for i, category in enumerate(['bachelor', 'master', 'staff', 'staff', 'external', 'student', 'doctoral', 'researcher']):
        r = Reader(full_name='Тест ' + str(i), ticket_number='T' + str(i), barcode='R' + str(i), category=category,
                   faculty='Ф1' if i in (0, 2) else 'Ф2' if i in (1, 3) else '',
                   department='К1' if i in (0, 2) else '', study_group='Г1' if i == 0 else '',
                   education_program='Программа' if i == 0 else '', course=1 if i == 0 else None,
                   cohort_year=2026 if i == 0 else None, is_teaching_staff=i == 2)
        s.add(r); people.append(r)
    s.flush()
    schedule = [
        (0, '2026-09-30T23:59', b1, 'desk'), (0, '2026-10-01T00:00', b1, 'desk'),
        (0, '2026-10-01T09:00', b1, 'desk'), (1, '2026-10-01T10:00', b2, 'kiosk'),
        (2, '2026-10-02T08:30', b1, 'desk'), (3, '2026-10-02T17:00', b2, 'desk'),
        (4, '2026-10-03T14:00', b2, 'desk'), (5, '2026-10-04T09:00', b1, 'desk'),
        (6, '2026-10-05T22:00', b1, 'turnstile'), (7, '2026-10-06T01:00', b1, 'desk'),
        (0, '2026-10-07T00:00', b1, 'desk'),
    ]
    for reader, stamp, branch, source in schedule:
        s.add(LibraryVisit(reader_id=people[reader].id, branch_id=branch.id, scanned_by=actor, source=source,
                           scanned_at=datetime.fromisoformat(stamp).replace(tzinfo=TZ).astimezone(timezone.utc)))
    s.commit()
    return b1.id, b2.id, actor, people[0].id

@pytest.fixture
def data(client):
    gen = app.dependency_overrides[db](); s = next(gen)
    try: ids = seed(s)
    finally: gen.close()
    return client, ids

def get(c, **kw):
    r = c.get(PATH, params={**PERIOD, **kw}); assert r.status_code == 200, r.text
    return r.json()

def test_totals_calendar_distinct_and_zero_days(data):
    c, _ = data
    r = get(c, date_from='2026-09-29')
    assert r['totals']['visits'] == 10 and r['totals']['unique'] == 8
    assert r['rows'][0]['visits'] == 0
    assert r['totals']['categories']['bachelor'] == {'visits': 3, 'unique': 1}
    assert r['totals']['categories']['faculty']['visits'] == 1
    assert r['totals']['categories']['staff']['visits'] == 1
    assert r['totals']['categories']['external']['visits'] == 1
    assert r['totals']['categories']['other']['visits'] == 2
    assert sum(v['visits'] for v in r['totals']['categories'].values()) == 10
    assert sum(row['visits'] for row in r['rows']) == 10
    assert sum(row['unique'] for row in r['months']) == 9  # Intentionally not the period's 8.
    assert r['summary']['peak'] == {'date': '2026-10-01', 'visits': 3}
    assert r['summary']['selected_days'] == 8 and r['summary']['active_days'] == 7
    assert get(c, date_from='2026-10-01', date_to='2026-10-01')['totals']['visits'] == 3
    weeks = get(c, group_by='week')
    assert [row['visits'] for row in weeks['rows']] == [8, 2]
    assert weeks['rows'][0]['date_from'] == '2026-09-30'
    months = get(c, group_by='month')
    assert [row['visits'] for row in months['rows']] == [1, 9]

def test_filters_combined_multi_values_and_empty(data):
    c, (b1, b2, actor, _) = data
    assert get(c, faculty=['Ф1', 'Ф2'])['totals']['visits'] == 6
    filters = dict(faculty=['Ф1'], department=['К1'], education_program=['Программа'],
                   study_group=['Г1'], course=['1'], cohort_year=['2026'], branch_id=[b1],
                   scanned_by=[actor], source=['desk'], bucket=['bachelor'], category=['bachelor'], status=['active'])
    assert get(c, **filters)['totals']['visits'] == 3
    assert get(c, faculty=[''], course=[''])['totals']['visits'] == 4
    assert get(c, branch_id=[b2])['totals']['visits'] == 3  # Closed branches still have history.
    assert get(c, weekday=[3, 4])['totals']['visits'] == 5
    assert get(c, time_from='09:00', time_to='17:00')['totals']['visits'] == 5
    assert get(c, time_from='22:00', time_to='08:00')['totals']['visits'] == 4
    assert get(c, q='T0')['totals']['visits'] == 3
    assert get(c, q='%')['totals']['visits'] == 0  # Literal wildcard, not match-all.
    assert get(c, faculty=['Несуществующий'])['totals']['unique'] == 0

def test_breakdowns_journal_and_options(data):
    c, (_, b2, _, _) = data
    r = get(c, breakdown='branch_id')
    assert sum(x['visits'] for x in r['breakdown']) == 10
    assert sum(x['visits'] for x in r['hours']) == 10
    assert sum(x['visits'] for x in r['weekdays']) == 10
    response = c.get(PATH + '/journal', params={**PERIOD, 'bucket': 'bachelor', 'limit': 2})
    assert response.status_code == 200
    j = response.json(); assert j['total'] == 3 and len(j['items']) == 2
    assert 'iin' not in str(j) and 'email' not in str(j)
    assert j['items'][0]['scanned_at'].endswith('+00:00')
    assert 'no-store' in response.headers['cache-control']
    assert c.get(PATH + '/journal', params={**PERIOD, 'bucket': 'bachelor', 'offset': 2}).json()['items'][0]['id'] != j['items'][0]['id']
    o = c.get(PATH + '/options').json()
    assert any(x['value'] == str(b2) and 'закрыт' in x['label'] for x in o['filters']['branch_id'])
    assert any(x['value'] == '' for x in o['filters']['faculty'])

@pytest.mark.parametrize('params', [
    {'date_from': '2026-10-07'}, {'date_to': '9999-12-31'}, {'date_from': '1900-01-01'},
    {'date_from': 'bad'}, {'time_from': '25:00'}, {'time_to': '8:30'}, {'weekday': ['7']},
    {'bucket': ['bad']}, {'branch_id': ['abc']}, {'source': ['bad']}, {'group_by': 'year'},
    {'q': 'x' * 201}, {'unknown_filter': 'bad'}, {'breakdown': 'iin'},
])
def test_invalid_filters(client, params):
    assert client.get(PATH, params={**PERIOD, **params}).status_code == 422

def test_exports_exact_filters_and_formula_escaping(data):
    c, _ = data
    params = {**PERIOD, 'faculty': 'Ф1'}
    response = c.get(PATH + '/export', params=params)
    assert response.status_code == 200 and response.content.startswith(b'PK')
    ws = load_workbook(io.BytesIO(response.content)).active
    total = next(r for r in ws.iter_rows(values_only=True) if r[0] == 'Итого за период')
    assert total[-2:] == (4, 2)
    assert ws.freeze_panes == 'B4'
    csv = c.get(PATH + '/export', params={**params, 'format': 'csv'}).text
    assert 'Итого за период' in csv and 'Ф1' in csv and 'Ф2' not in csv
    from app.attendance import Filters, export_report, report
    gen = app.dependency_overrides[db](); s = next(gen)
    try:
        reader = s.scalar(select(Reader).where(Reader.category == 'external')); reader.faculty = '=1+1'; s.commit()
        f = Filters(**PERIOD)
        response = export_report(report(s, f), f, 'xlsx')
        ws = load_workbook(io.BytesIO(response.body)).active
        assert any(cell.value == "'=1+1" and cell.data_type == 's' for row in ws for cell in row)
        assert not any(cell.data_type == 'f' for row in ws for cell in row)
    finally: gen.close()

def test_permissions_and_external_reader(client):
    r = client.post('/api/readers', json={'full_name': 'Внешний читатель', 'category': 'external'})
    assert r.status_code == 200 and '-E' in r.json()['ticket_number']
    add(client, permissions=['visits.view'])
    with TestClient(app) as c:
        assert c.get(PATH).status_code == 401
        login(c)
        for suffix in ('', '/options', '/journal', '/export'):
            assert c.get(PATH + suffix, params=PERIOD).status_code == 200
        assert c.get('/api/statistics/staff').status_code == 403
        assert c.post('/api/visits', json={}).status_code == 403
    add(client, username='stats', permissions=['statistics.view'])
    with TestClient(app) as c:
        login(c, name='stats')
        assert c.get(PATH).status_code == 403

def test_private_search_not_logged():
    import logging
    from app.readers import ReaderAccessLogFilter
    record = logging.LogRecord('uvicorn.access', 20, '', 0, '%s %s %s %s %s',
                               ('local', 'GET', PATH + '?q=private-name', '1.1', 200), None)
    ReaderAccessLogFilter().filter(record)
    assert 'private-name' not in record.getMessage()
